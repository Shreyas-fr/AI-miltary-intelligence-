import io

def parse_pdf(file_bytes: bytes) -> str:
    import pdfplumber
    text = ""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

def parse_docx(file_bytes: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join([para.text for para in doc.paragraphs])

def parse_xlsx(file_bytes: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    text = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_vals = [str(val) for val in row if val is not None]
            if row_vals:
                text.append(" | ".join(row_vals))
    return "\n".join(text)

def parse_file(filename: str, file_bytes: bytes) -> str:
    """Parses a file's bytes into text based on its extension."""
    ext = filename.lower().split('.')[-1]
    
    if ext == "pdf":
        return parse_pdf(file_bytes)
    elif ext == "docx":
        return parse_docx(file_bytes)
    elif ext == "xlsx":
        return parse_xlsx(file_bytes)
    elif ext in ["txt", "md", "csv"]:
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {ext}")
